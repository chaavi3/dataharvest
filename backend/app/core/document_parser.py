"""Parse uploaded documents into plain text for LLM extraction."""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


async def parse_document(file_path: str) -> str:
    p = Path(file_path)
    suffix = p.suffix.lower()

    try:
        if suffix == ".txt":
            return p.read_text(encoding="utf-8", errors="replace")

        if suffix == ".md":
            return p.read_text(encoding="utf-8", errors="replace")

        if suffix == ".csv":
            return p.read_text(encoding="utf-8", errors="replace")

        if suffix == ".html" or suffix == ".htm":
            return _parse_html(p)

        if suffix == ".pdf":
            return _parse_pdf(p)

        if suffix == ".docx":
            return _parse_docx(p)

        if suffix in (".xlsx", ".xls"):
            return _parse_xlsx(p)

        logger.warning("Unsupported file type: %s", suffix)
        return p.read_text(encoding="utf-8", errors="replace")

    except Exception as exc:
        logger.error("Failed to parse %s: %s", file_path, exc)
        return ""


def _parse_pdf(p: Path) -> str:
    try:
        import pdfplumber
        text_parts: list[str] = []
        with pdfplumber.open(str(p)) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
        return "\n\n".join(text_parts)
    except ImportError:
        logger.warning("pdfplumber not installed; trying PyPDF2")
        from PyPDF2 import PdfReader
        reader = PdfReader(str(p))
        return "\n\n".join(pg.extract_text() or "" for pg in reader.pages)


def _parse_docx(p: Path) -> str:
    from docx import Document
    doc = Document(str(p))
    return "\n\n".join(para.text for para in doc.paragraphs if para.text.strip())


def _parse_xlsx(p: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        rows_text: list[str] = []
        for row in ws.iter_rows(values_only=True):
            rows_text.append("\t".join(str(c) if c is not None else "" for c in row))
        parts.append("\n".join(rows_text))
    return "\n\n---\n\n".join(parts)


def _parse_html(p: Path) -> str:
    import re
    raw = p.read_text(encoding="utf-8", errors="replace")
    text = re.sub(r"<script[^>]*>.*?</script>", "", raw, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<style[^>]*>.*?</style>", "", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text
